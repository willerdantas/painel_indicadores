"""
importacao_custos.py

Importação da planilha mensal do setor de custos (padrão "Planilha-do-mes-de-
<Mês>-<Ano>.xlsx"), só a aba GERAL:

- linhas: itens de custo agrupados nas seções da planilha (1.1 Pessoal,
  1.2 Material de Consumo, 1.3 Despesas Gerais, 2 Custos Indiretos), o rateio
  dos centros auxiliares (6.x) e a produção de cada unidade (8);
- colunas: D = total do hospital; da E em diante, uma por unidade (centro de
  custo), agrupadas pela linha acima do cabeçalho (células mescladas).
Subtotais, totais, custo gerencial, custo com rateio e custo unitário são
contas da própria planilha e não são importados.

Fluxo: preparar() lê a planilha e grava num lote de preparação
(importacao_custos_lote / _linha); a pré-visualização é montada a partir do
lote (carregar()); confirmar() grava tudo numa transação, substituindo o mês
em custos_unidade, rateio_unidade e producao_unidade, refaz o total por
despesa em custos (soma das unidades) e apaga o lote.
Itens, unidades, grupos e itens de rateio que ainda não existem são
cadastrados na confirmação (a pré-visualização os marca como novos).
"""

import json
import re
import unicodedata
import uuid
from decimal import Decimal

import openpyxl
from psycopg2.extras import execute_batch

import banco_painel
import lancamento_custos
from banco_painel import _cursor, schema

# seção da planilha -> tipo de custo (os três primeiros já existem em tipo_custo)
TIPOS = {'1.1': 'Pessoal', '1.2': 'Material de Consumo', '1.3': 'Despesas Gerais', '2': 'Custos Indiretos'}
_SECAO_CUSTO_TOTAL = '5'
_SECAO_RATEIO = '6'
_SECAO_PRODUCAO = '8'
_SECOES_CALCULADAS = ('3', '4', '7', '9')  # despesas gerais, custo gerencial, custo c/ rateio, custo unitário
_CODIGO = re.compile(r'^\s*(\d+(?:\.\d+)*)\.?\s*[-–—]?\s*(.*)$')
_TOLERANCIA = 0.05  # R$: arredondamento entre o total (coluna D) e a soma das unidades
_DIAS_LOTE = 2      # lotes não confirmados são apagados depois disso
TABELAS_DO_MES = ('custos', 'custos_unidade', 'rateio_unidade', 'producao_unidade')


class ErroPlanilha(ValueError):
    """Planilha fora do padrão esperado."""


def _limpar(texto):
    return ' '.join(str(texto).split()) if texto is not None else ''


def _chave(texto):
    """Nome para comparar: sem acentos, maiúsculas e espaços repetidos."""
    sem_acento = unicodedata.normalize('NFKD', _limpar(texto)).encode('ascii', 'ignore').decode()
    return sem_acento.upper()


def _moeda(v):
    return f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _numero(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return banco_painel.ler_numero(v)  # '1.234,56'; '%', '#DIV/0!', 'Valor (R$)' -> None
    except ValueError:
        return None


def _aba_geral(wb):
    for ws in wb.worksheets:
        if _chave(ws.title) == 'GERAL':
            return ws
    raise ErroPlanilha('a planilha não tem a aba GERAL (abas: ' + ', '.join(ws.title.strip() for ws in wb.worksheets) + ').')


def _linha_cabecalho(ws):
    """Linha com 'Item' na coluna B e 'valor' na D (as unidades vêm a partir da E)."""
    for r in range(1, min(ws.max_row, 60) + 1):
        if _chave(ws.cell(r, 2).value) == 'ITEM' and _chave(ws.cell(r, 4).value) == 'VALOR':
            return r
    raise ErroPlanilha('cabeçalho não encontrado na aba GERAL (linha com "Item" na coluna B e "valor" na coluna D).')


def _unidades(ws, cab):
    """Colunas com nome no cabeçalho, da E em diante; o grupo vem da linha de
    cima (mesclada sobre as unidades do grupo)."""
    colunas = [c for c in range(5, ws.max_column + 1) if _limpar(ws.cell(cab, c).value)]
    if not colunas:
        raise ErroPlanilha('nenhuma unidade (centro de custo) no cabeçalho, da coluna E em diante.')
    grupo_da_coluna = {}
    for faixa in ws.merged_cells.ranges:
        if faixa.min_row <= cab - 1 <= faixa.max_row:
            nome = _limpar(ws.cell(faixa.min_row, faixa.min_col).value)
            for c in range(faixa.min_col, faixa.max_col + 1):
                grupo_da_coluna[c] = nome
    grupo_atual, unidades = None, []
    for c in colunas:
        grupo_atual = grupo_da_coluna.get(c) or _limpar(ws.cell(cab - 1, c).value) or grupo_atual
        if not grupo_atual:
            raise ErroPlanilha(f'a unidade da coluna {ws.cell(cab, c).column_letter} não tem grupo na linha {cab - 1}.')
        unidades.append({'coluna': c, 'nome': _limpar(ws.cell(cab, c).value), 'grupo': grupo_atual,
                         'medida': None, 'producao': None})
    nomes = [_chave(u['nome']) for u in unidades]
    repetidas = sorted({u['nome'] for u in unidades if nomes.count(_chave(u['nome'])) > 1})
    if repetidas:
        raise ErroPlanilha('unidades repetidas no cabeçalho: ' + ', '.join(repetidas) + '.')
    return unidades


def _secao(codigo):
    partes = codigo.split('.')
    return TIPOS.get('.'.join(partes[:2])) or TIPOS.get(partes[0])


def ler_geral(caminho):
    """Lê a aba GERAL. Retorna {mes_planilha, unidades, itens, avisos,
    custo_total_planilha}; levanta ErroPlanilha se estiver fora do padrão."""
    try:
        wb = openpyxl.load_workbook(caminho, data_only=True)
    except Exception as e:
        raise ErroPlanilha(f'não foi possível abrir a planilha: {e}') from e
    ws = _aba_geral(wb)
    cab = _linha_cabecalho(ws)

    mes_planilha = None
    for linha in ws.iter_rows(min_row=1, max_row=cab - 1):
        for c in linha:
            if isinstance(c.value, str) and _chave(c.value).startswith('MES'):
                mes_planilha = _limpar(c.value.split(':', 1)[-1])
    unidades = _unidades(ws, cab)

    itens, avisos = [], []
    secao = None  # tipo de custo ou 'rateio'
    custo_total_planilha = None
    depois_producao = False
    for r in range(cab + 1, ws.max_row + 1):
        rotulo = ws.cell(r, 1).value
        valores = {u['nome']: ws.cell(r, u['coluna']).value for u in unidades}

        if not _limpar(rotulo):
            # medida da produção (Diária, Atendimento...): linha sem rótulo logo abaixo da produção
            textos = {k: _limpar(v) for k, v in valores.items() if isinstance(v, str) and v.strip()}
            if depois_producao and textos:
                for u in unidades:
                    u['medida'] = textos.get(u['nome'], u['medida'])
                depois_producao = False
            continue

        m = _CODIGO.match(str(rotulo))
        codigo, nome = (m.group(1), _limpar(m.group(2))) if m else (None, _limpar(rotulo))
        primeiro = codigo.split('.')[0] if codigo else None
        total = _numero(ws.cell(r, 4).value)
        numeros = {k: _numero(v) for k, v in valores.items()}

        if primeiro == _SECAO_PRODUCAO:
            for u in unidades:
                u['producao'] = numeros[u['nome']]
            depois_producao = True
            continue
        if primeiro == _SECAO_CUSTO_TOTAL:
            custo_total_planilha = total
            continue
        if primeiro in _SECOES_CALCULADAS or 'TOTAL' in _chave(nome):
            continue  # subtotais, totais e contas da própria planilha
        if total is None and all(v is None for v in numeros.values()):
            # título de seção (1.2 – MATERIAL DE CONSUMO, 6 Rateio...)
            if primeiro == _SECAO_RATEIO:
                secao = 'rateio'
            elif codigo:
                secao = _secao(codigo) or secao
            continue

        if primeiro == _SECAO_RATEIO:
            secao = 'rateio'
        elif codigo:
            secao = _secao(codigo)
        if not secao:
            avisos.append(f'Linha {r} ("{nome}") ignorada: fora das seções de custo conhecidas.')
            continue
        item = {'linha': r, 'categoria': 'rateio' if secao == 'rateio' else 'custo',
                'tipo': None if secao == 'rateio' else secao, 'item': nome, 'total': total,
                'valores': {k: v for k, v in numeros.items() if v is not None}}
        itens.append(item)
        if item['categoria'] == 'custo' and total is not None:
            soma = sum(item['valores'].values())
            if abs(total - soma) > _TOLERANCIA:
                avisos.append(f'Linha {r} ("{nome}"): total da coluna D ({_moeda(total)}) diferente da '
                              f'soma das unidades ({_moeda(soma)}).')

    if not any(i['categoria'] == 'custo' for i in itens):
        raise ErroPlanilha('nenhum item de custo encontrado na aba GERAL.')
    vistos = set()
    for i in itens:
        k = (i['categoria'], i['tipo'], _chave(i['item']))
        if k in vistos:
            raise ErroPlanilha(f'item repetido na linha {i["linha"]}: "{i["item"]}".')
        vistos.add(k)
    if custo_total_planilha is not None:
        soma = sum(i['total'] or 0 for i in itens if i['categoria'] == 'custo')
        if abs(soma - custo_total_planilha) > 1:
            avisos.append(f'A soma dos itens ({_moeda(soma)}) não bate com o CUSTO TOTAL da planilha '
                          f'({_moeda(custo_total_planilha)}).')
    return {'mes_planilha': mes_planilha, 'unidades': unidades, 'itens': itens, 'avisos': avisos,
            'custo_total_planilha': custo_total_planilha}


# ---------------------------------------------------------------- preparação

def preparar(caminho, arquivo, mes, usuario):
    """Lê a planilha e grava no lote de preparação. Retorna o id do lote."""
    dados = ler_geral(caminho)
    lote = uuid.uuid4().hex
    s = schema()
    coluna = {u['nome']: u['coluna'] for u in dados['unidades']}
    grupo = {u['nome']: u['grupo'] for u in dados['unidades']}
    # producao: uma linha por unidade (leva grupo e medida, mesmo sem produção);
    # custo/rateio: o total (unidade vazia) e uma linha por unidade com valor
    linhas = [['producao', 0, u['coluna'], None, None, u['grupo'], u['nome'], u['medida'], u['producao']]
              for u in dados['unidades']]
    for i in dados['itens']:
        linhas.append([i['categoria'], i['linha'], None, i['tipo'], i['item'], None, None, None, i['total']])
        linhas += [[i['categoria'], i['linha'], coluna[nome], i['tipo'], i['item'], grupo[nome], nome, None, valor]
                   for nome, valor in i['valores'].items()]
    with _cursor() as cur:
        cur.execute(f"DELETE FROM {s}.importacao_custos_lote WHERE criado_em < now() - interval '{_DIAS_LOTE} days'")
        cur.execute(f'INSERT INTO {s}.importacao_custos_lote (id, arquivo, mes, mes_planilha, custo_total, avisos, usuario) '
                    'VALUES (%s, %s, %s, %s, %s, %s, %s)',
                    [lote, arquivo, f'{mes}-01', dados['mes_planilha'], dados['custo_total_planilha'],
                     json.dumps(dados['avisos'], ensure_ascii=False), usuario])
        execute_batch(cur, f"""
            INSERT INTO {s}.importacao_custos_linha
                (lote_id, categoria, linha, coluna, tipo, item, grupo_unidade, unidade, medida, valor)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""", [[lote] + l for l in linhas], page_size=500)
    return lote


def cancelar(lote):
    with _cursor() as cur:
        cur.execute(f'DELETE FROM {schema()}.importacao_custos_lote WHERE id = %s', [lote])


def _cadastros(cur):
    """Cadastros atuais por nome normalizado (item_custo por (tipo, item))."""
    s = schema()
    cur.execute(f'SELECT id, nome FROM {s}.tipo_custo')
    tipos = {_chave(r['nome']): r for r in cur.fetchall()}
    cur.execute(f'SELECT it.id, it.nome, t.nome AS tipo FROM {s}.item_custo it '
                f'JOIN {s}.tipo_custo t ON t.id = it.tipo_custo_id')
    itens = {(_chave(r['tipo']), _chave(r['nome'])): r for r in cur.fetchall()}
    cur.execute(f'SELECT id, nome FROM {s}.grupo_unidade')
    grupos = {_chave(r['nome']): r for r in cur.fetchall()}
    cur.execute(f'SELECT id, nome, medida FROM {s}.unidade')
    unidades = {_chave(r['nome']): r for r in cur.fetchall()}
    cur.execute(f'SELECT id, nome FROM {s}.item_rateio')
    rateios = {_chave(r['nome']): r for r in cur.fetchall()}
    return {'tipo_custo': tipos, 'item_custo': itens, 'grupo_unidade': grupos, 'unidade': unidades,
            'item_rateio': rateios}


def _ler_lote(cur, lote):
    s = schema()
    cur.execute(f'SELECT * FROM {s}.importacao_custos_lote WHERE id = %s', [lote])
    cab = cur.fetchone()
    if cab is None:
        return None
    cur.execute(f'SELECT * FROM {s}.importacao_custos_linha WHERE lote_id = %s ORDER BY linha, coluna NULLS FIRST', [lote])
    linhas = [{k: float(v) if isinstance(v, Decimal) else v for k, v in r.items()} for r in cur.fetchall()]

    unidades = [{'nome': l['unidade'], 'grupo': l['grupo_unidade'], 'medida': l['medida'],
                 'producao': l['valor'], 'coluna': l['coluna']}
                for l in sorted((l for l in linhas if l['categoria'] == 'producao'), key=lambda l: l['coluna'])]
    itens, por_chave = [], {}
    for l in linhas:
        if l['categoria'] == 'producao':
            continue
        k = (l['categoria'], l['tipo'], l['item'])
        if k not in por_chave:
            por_chave[k] = {'categoria': l['categoria'], 'tipo': l['tipo'], 'item': l['item'], 'linha': l['linha'],
                            'total': None, 'valores': {}}
            itens.append(por_chave[k])
        if l['unidade'] is None:
            por_chave[k]['total'] = l['valor']
        else:
            por_chave[k]['valores'][l['unidade']] = l['valor']
    return {'id': cab['id'], 'arquivo': cab['arquivo'], 'mes': cab['mes'].strftime('%Y-%m'),
            'mes_planilha': cab['mes_planilha'], 'usuario': cab['usuario'], 'criado_em': cab['criado_em'],
            'custo_total_planilha': float(cab['custo_total']) if cab['custo_total'] is not None else None,
            'avisos': json.loads(cab['avisos']), 'unidades': unidades, 'itens': itens}


_MESES = ['JANEIRO', 'FEVEREIRO', 'MARCO', 'ABRIL', 'MAIO', 'JUNHO', 'JULHO', 'AGOSTO', 'SETEMBRO',
          'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']


def _mes_divergente(mes_planilha, mes):
    """True se o "MÊS:" escrito na planilha não é o mês escolhido para o
    lançamento (o cabeçalho às vezes fica desatualizado; vale o escolhido)."""
    if not mes_planilha:
        return False
    texto = _chave(mes_planilha)
    ano, numero = mes.split('-')
    return _MESES[int(numero) - 1] not in texto or ano not in texto


def carregar(lote, mes=None):
    """Lote de preparação montado para a pré-visualização (None se não existir):
    o que é novo nos cadastros e o que já existe no mês e será substituído."""
    s = schema()
    with _cursor() as cur:
        dados = _ler_lote(cur, lote)
        if dados is None:
            return None
        mes = mes or dados['mes']
        cad = _cadastros(cur)
        existentes = {}
        for tabela in TABELAS_DO_MES:
            cur.execute(f'SELECT COUNT(*) AS n FROM {s}.{tabela} WHERE mes = %s', [f'{mes}-01'])
            existentes[tabela] = cur.fetchone()['n']

    custos = [i for i in dados['itens'] if i['categoria'] == 'custo']
    rateios = [i for i in dados['itens'] if i['categoria'] == 'rateio']
    for i in custos:
        i['novo'] = (_chave(i['tipo']), _chave(i['item'])) not in cad['item_custo']
        i['soma'] = sum(i['valores'].values())
    for i in rateios:
        i['novo'] = _chave(i['item']) not in cad['item_rateio']
    grupos = []
    for u in dados['unidades']:
        u['novo'] = _chave(u['nome']) not in cad['unidade']
        u['custo'] = sum(i['valores'].get(u['nome'], 0) for i in custos)
        if not grupos or grupos[-1]['nome'] != u['grupo']:
            grupos.append({'nome': u['grupo'], 'novo': _chave(u['grupo']) not in cad['grupo_unidade'], 'unidades': []})
        grupos[-1]['unidades'].append(u)
    tipos = []
    for i in custos:
        if not tipos or tipos[-1]['nome'] != i['tipo']:
            tipos.append({'nome': i['tipo'], 'novo': _chave(i['tipo']) not in cad['tipo_custo'], 'itens': [], 'total': 0})
        tipos[-1]['itens'].append(i)
        tipos[-1]['total'] += i['total'] or 0
    dados.update(
        mes=mes, mes_divergente=_mes_divergente(dados['mes_planilha'], mes), grupos=grupos, tipos=tipos, custos=custos, rateios=rateios, existentes=existentes,
        total_custos=sum(t['total'] for t in tipos),
        novos={'tipos': sum(t['novo'] for t in tipos), 'itens': sum(i['novo'] for i in custos),
               'grupos': sum(g['novo'] for g in grupos), 'unidades': sum(u['novo'] for u in dados['unidades']),
               'rateios': sum(i['novo'] for i in rateios)},
    )
    return dados


# ---------------------------------------------------------------- confirmação

def aplicar(cur, lote, mes, usuario):
    """Grava o lote no mês (substituindo o que houver nele) e apaga o lote, no
    cursor dado: a transação é de quem chama. Retorna o resumo do que gravou."""
    s = schema()
    dados = _ler_lote(cur, lote)
    if dados is None:
        raise ErroPlanilha('esta importação não existe mais (já foi confirmada ou cancelada).')
    cad = _cadastros(cur)
    data = f'{mes}-01'

    def obter_id(tabela, chave, sql, params):
        """id do cadastro; cria se ainda não existir."""
        if chave not in cad[tabela]:
            cur.execute(sql, params)
            cad[tabela][chave] = cur.fetchone()
        return cad[tabela][chave]['id']

    def tipo_id(nome):
        return obter_id('tipo_custo', _chave(nome), f'INSERT INTO {s}.tipo_custo (nome) VALUES (%s) RETURNING id', [nome])

    def item_id(tipo, nome):
        return obter_id('item_custo', (_chave(tipo), _chave(nome)),
                        f'INSERT INTO {s}.item_custo (tipo_custo_id, nome) VALUES (%s, %s) RETURNING id',
                        [tipo_id(tipo), nome])

    def grupo_id(nome):
        return obter_id('grupo_unidade', _chave(nome),
                        f'INSERT INTO {s}.grupo_unidade (nome) VALUES (%s) RETURNING id', [nome])

    def rateio_id(nome):
        return obter_id('item_rateio', _chave(nome),
                        f'INSERT INTO {s}.item_rateio (nome) VALUES (%s) RETURNING id', [nome])

    unidade_id = {}
    for u in dados['unidades']:
        unidade_id[u['nome']] = obter_id(
            'unidade', _chave(u['nome']),
            f'INSERT INTO {s}.unidade (grupo_unidade_id, nome, medida) VALUES (%s, %s, %s) RETURNING id',
            [grupo_id(u['grupo']), u['nome'], u['medida']])
        if u['medida']:  # unidade já cadastrada sem medida: completa com a da planilha
            cur.execute(f'UPDATE {s}.unidade SET medida = %s WHERE id = %s AND medida IS NULL',
                        [u['medida'], unidade_id[u['nome']]])

    for tabela in TABELAS_DO_MES:
        cur.execute(f'DELETE FROM {s}.{tabela} WHERE mes = %s', [data])

    por_unidade, rateio = [], []
    for i in dados['itens']:
        if i['categoria'] == 'custo':
            iid = item_id(i['tipo'], i['item'])
            por_unidade += [[data, iid, unidade_id[u], v, usuario] for u, v in i['valores'].items()]
        else:
            rid = rateio_id(i['item'])
            rateio += [[data, rid, unidade_id[u], v, usuario] for u, v in i['valores'].items()]
    producao = [[data, unidade_id[u['nome']], u['producao'], usuario]
                for u in dados['unidades'] if u['producao'] is not None]

    for sql, params in (
        (f'INSERT INTO {s}.custos_unidade (mes, item_custo_id, unidade_id, valor, atualizado_por) '
         'VALUES (%s, %s, %s, %s, %s)', por_unidade),
        (f'INSERT INTO {s}.rateio_unidade (mes, item_rateio_id, unidade_id, valor, atualizado_por) '
         'VALUES (%s, %s, %s, %s, %s)', rateio),
        (f'INSERT INTO {s}.producao_unidade (mes, unidade_id, producao, atualizado_por) '
         'VALUES (%s, %s, %s, %s)', producao),
    ):
        execute_batch(cur, sql, params, page_size=500)
    despesas = lancamento_custos.consolidar(cur, mes)  # total por despesa = soma das unidades
    cur.execute(f'DELETE FROM {s}.importacao_custos_lote WHERE id = %s', [lote])
    return {'itens': despesas, 'valores_unidade': len(por_unidade), 'rateio': len(rateio),
            'unidades_com_producao': len(producao)}


def confirmar(lote, mes, usuario):
    with _cursor() as cur:  # tudo ou nada
        return aplicar(cur, lote, mes, usuario)
