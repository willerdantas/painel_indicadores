"""
financeiro.py

Página pública "Financeiro" e a importação da tela "Lançar Dados do
Financeiro": dois relatórios do SIAFE (Secretaria da Fazenda), em .xls ou
.xlsx, primeira aba.

- Empenhos em aberto ("Relatório de Empenhos em aberto por fonte de recursos"):
  linhas 1 a 8 de cabeçalho; A = empenho, B = credor ("CNPJ - nome"),
  C = empenhos a liquidar, D = empenhos liquidados a pagar. Os empenhos vêm
  agrupados por fonte de recurso (linha "500 - RECURSOS ..." acima deles).
  O relatório só detalha os empenhos com valor a liquidar: o total de
  liquidados a pagar de cada fonte (que inclui os não detalhados) é guardado
  à parte, em totais da importação.
- Restos a pagar ("Restos a Pagar Inscritos, Pagos e A Pagar - Nota de
  Empenho e Ordem Bancária OB"): linhas 1 a 7 de cabeçalho; uma linha por
  ordem bancária, com empenho, processo, fonte, contrato e credor em células
  mescladas (repetidos em todas as linhas da mescla). As colunas são achadas
  pelo nome no cabeçalho (linha 6).

O processo perde "NUP", espaços, "-" e "."; o credor é separado em documento
(CNPJ/CPF, só dígitos) e nome, guardados no cadastro credor. Cada importação
substitui a anterior do mesmo relatório (fica só a posição mais recente). Os
dois relatórios se cruzam pelo número do empenho.

Tabelas (schema do painel, criadas em banco_painel.iniciar):
  credor (documento, nome)
  importacao_financeiro (id, tipo, arquivo, referencia, atualizado_siafe,
                         linhas, totais, importado_em, importado_por)
  empenho_aberto (importacao_id, empenho, fonte, credor_documento,
                  valor_a_liquidar, valor_liquidado_a_pagar)
  resto_pagar (importacao_id, empenho, processo, fonte, contrato,
               credor_documento, ordem_bancaria, data_ob, valor_inscrito,
               valor_pago, valor_a_pagar)
"""

import io
import json
import os
import re
import unicodedata
from datetime import date, datetime

from psycopg2.extras import execute_values

import banco_painel
from banco_painel import _cursor, schema

TIPOS = {
    'empenhos': 'Empenhos em aberto',
    'restos': 'Restos a pagar',
}
EXTENSOES = ('.xls', '.xlsx')

_EMPENHO = re.compile(r'^\d{4}NE\d+$')
_ORDEM_BANCARIA = re.compile(r'^\d{4}OB\d+$')
_FONTE = re.compile(r'^\d+ - ')
# linhas de cabeçalho de cada relatório (os dados começam na seguinte)
_LINHAS_CABECALHO = {'empenhos': 8, 'restos': 7}
_LINHA_TITULO = 3          # linha 4: título com o período
_LINHA_COLUNAS = 5         # linha 6: nomes das colunas
_LINHA_TOTAL = 6           # linha 7: total da unidade gestora
_TOLERANCIA = 0.01         # conferência dos totais (centavos)


class PlanilhaInvalida(ValueError):
    """Arquivo que não é o relatório esperado."""


# ---------------------------------------------------------------- leitura

def _grade(caminho):
    """Valores da primeira aba (lista de linhas), com o valor de cada célula
    mesclada repetido em todas as células da mescla. Datas viram datetime."""
    if caminho.lower().endswith('.xls'):
        import xlrd
        livro = xlrd.open_workbook(caminho, formatting_info=True)
        aba = livro.sheet_by_index(0)

        def valor(r, c):
            if aba.cell_type(r, c) == xlrd.XL_CELL_DATE:
                return xlrd.xldate_as_datetime(aba.cell_value(r, c), livro.datemode)
            return aba.cell_value(r, c)
        grade = [[valor(r, c) for c in range(aba.ncols)] for r in range(aba.nrows)]
        mesclas = aba.merged_cells  # (linha ini, linha fim, coluna ini, coluna fim), fim exclusivo
    else:
        import openpyxl
        aba = openpyxl.load_workbook(caminho, data_only=True).worksheets[0]
        grade = [['' if c.value is None else c.value for c in linha] for linha in aba.iter_rows()]
        mesclas = [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col) for m in aba.merged_cells.ranges]
    for r0, r1, c0, c1 in mesclas:
        for r in range(r0, min(r1, len(grade))):
            for c in range(c0, min(c1, len(grade[r]))):
                grade[r][c] = grade[r0][c0]
    return grade


def _texto(v):
    return '' if v is None else str(v).strip()


def _numero(v):
    return banco_painel.ler_numero(v) or 0.0


def _celula(grade, linha, coluna):
    return grade[linha][coluna] if linha < len(grade) and coluna < len(grade[linha]) else ''


def _sem_acento(texto):
    return unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode().lower()


def separar_credor(texto):
    """("54858014000927", "ESSITY DO BRASIL") de "54858014000927 - ESSITY DO
    BRASIL": documento só com dígitos; sem " - ", tudo é o nome."""
    texto = _texto(texto)
    documento, separador, nome = texto.partition(' - ')
    if not separador:
        return '', texto
    return re.sub(r'\D', '', documento), nome.strip()


def limpar_processo(texto):
    """"NUP 24001.006317/2024-93" -> "24001006317/202493" (sem NUP, espaços, - e .)."""
    return re.sub(r'NUP|[\s.\-]', '', _texto(texto), flags=re.IGNORECASE) or None


def _data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(_texto(v), '%d/%m/%Y').date()
    except ValueError:
        return None


def _referencia(grade):
    """Período do título (linha 4): "Até 12 de 2025", "Até o Mês 10 / 2026"."""
    titulo = _texto(_celula(grade, _LINHA_TITULO, 1)) or _texto(_celula(grade, _LINHA_TITULO, 0))
    m = re.search(r'(At[ée]\b.*)$', titulo)
    return (m.group(1) if m else titulo).strip() or None


def _atualizado_siafe(grade):
    """"Dados atualizados em: 23/10/2026 10:41:05" do rodapé do relatório."""
    for linha in grade:
        m = re.search(r'Dados atualizados em:\s*(.+)$', _texto(linha[0]) if linha else '')
        if m:
            return m.group(1).strip()
    return None


def _confere(soma, total):
    return abs(round(soma, 2) - round(total, 2)) <= _TOLERANCIA


def ler_empenhos(caminho):
    """Relatório de empenhos em aberto -> {linhas, referencia, atualizado_siafe, totais}."""
    grade = _grade(caminho)
    colunas = [_sem_acento(_texto(_celula(grade, _LINHA_COLUNAS, c))) for c in range(4)]
    if 'liquidar' not in colunas[2] or 'liquidad' not in colunas[3]:
        raise PlanilhaInvalida('não é o "Relatório de Empenhos em aberto por fonte de recursos" '
                               '(na linha 6, as colunas C e D devem ser "Empenhos a Liquidar" e '
                               '"Empenhos Liquidadas a Pagar").')
    linhas, fontes, fonte = [], {}, None
    # a linha da fonte vem antes dos empenhos dela (a primeira fica na linha 8, do cabeçalho)
    for r in range(_LINHA_TOTAL + 1, len(grade)):
        a = _texto(_celula(grade, r, 0))
        if _EMPENHO.match(a):
            if r < _LINHAS_CABECALHO['empenhos']:
                continue
            documento, nome = separar_credor(_celula(grade, r, 1))
            linhas.append({'empenho': a, 'fonte': fonte, 'credor_documento': documento, 'credor_nome': nome,
                           'valor_a_liquidar': _numero(_celula(grade, r, 2)),
                           'valor_liquidado_a_pagar': _numero(_celula(grade, r, 3))})
        elif _FONTE.match(a) and _texto(_celula(grade, r, 1)) in ('', a):  # B vazia ou mesclada com A
            fonte = a
            fontes[fonte] = {'valor_a_liquidar': _numero(_celula(grade, r, 2)),
                             'valor_liquidado_a_pagar': _numero(_celula(grade, r, 3))}
    if not linhas:
        raise PlanilhaInvalida('nenhum empenho encontrado (coluna A no formato 2025NE000001).')
    total = {'valor_a_liquidar': _numero(_celula(grade, _LINHA_TOTAL, 2)),
             'valor_liquidado_a_pagar': _numero(_celula(grade, _LINHA_TOTAL, 3))}
    soma = {k: round(sum(l[k] for l in linhas), 2) for k in total}
    return {'linhas': linhas, 'referencia': _referencia(grade), 'atualizado_siafe': _atualizado_siafe(grade),
            'totais': {'relatorio': total, 'detalhado': soma, 'fontes': fontes,
                       'confere': {k: _confere(soma[k], total[k]) for k in total}}}


# colunas dos restos a pagar: nome no cabeçalho (sem acento, minúsculas) -> campo
_COLUNAS_RESTOS = [
    ('nota de empenho', 'empenho'), ('processo', 'processo'), ('fonte', 'fonte'),
    ('contrato', 'contrato'), ('credor', 'credor'), ('ordem bancaria', 'ordem_bancaria'),
    ('data emissao ob', 'data_ob'), ('total inscrito', 'valor_inscrito'),
    ('valor pago', 'valor_pago'), ('valor a pagar', 'valor_a_pagar'),
]


def ler_restos(caminho):
    """Relatório de restos a pagar -> {linhas, referencia, atualizado_siafe, totais}."""
    grade = _grade(caminho)
    cabecalho = [_sem_acento(_texto(c)) for c in (grade[_LINHA_COLUNAS] if len(grade) > _LINHA_COLUNAS else [])]
    col = {}
    for nome, campo in _COLUNAS_RESTOS:
        col[campo] = next((i for i, c in enumerate(cabecalho) if nome in c), None)
    faltando = [nome for nome, campo in _COLUNAS_RESTOS if col[campo] is None]
    if faltando:
        raise PlanilhaInvalida('não é o relatório "Restos a Pagar Inscritos, Pagos e A Pagar": na linha 6 '
                               'faltam as colunas ' + ', '.join(f'"{n}"' for n in faltando) + '.')
    linhas = []
    for r in range(_LINHAS_CABECALHO['restos'], len(grade)):
        empenho = _texto(_celula(grade, r, col['empenho']))
        if not _EMPENHO.match(empenho):
            continue  # total da unidade gestora, rodapé e filtros do relatório
        documento, nome = separar_credor(_celula(grade, r, col['credor']))
        ob = _texto(_celula(grade, r, col['ordem_bancaria']))
        linhas.append({
            'empenho': empenho,
            'processo': limpar_processo(_celula(grade, r, col['processo'])),
            'fonte': _texto(_celula(grade, r, col['fonte'])) or None,
            'contrato': _texto(_celula(grade, r, col['contrato'])) or None,
            'credor_documento': documento, 'credor_nome': nome,
            'ordem_bancaria': ob if _ORDEM_BANCARIA.match(ob) else None,
            'data_ob': _data(_celula(grade, r, col['data_ob'])) if _ORDEM_BANCARIA.match(ob) else None,
            'valor_inscrito': _numero(_celula(grade, r, col['valor_inscrito'])),
            'valor_pago': _numero(_celula(grade, r, col['valor_pago'])),
            'valor_a_pagar': _numero(_celula(grade, r, col['valor_a_pagar'])),
        })
    if not linhas:
        raise PlanilhaInvalida('nenhum empenho encontrado (coluna da nota de empenho no formato 2024NE000001).')
    campos = ('valor_inscrito', 'valor_pago', 'valor_a_pagar')
    total = {k: _numero(_celula(grade, _LINHA_TOTAL, col[k])) for k in campos}
    soma = {k: round(sum(l[k] for l in linhas), 2) for k in campos}
    return {'linhas': linhas, 'referencia': _referencia(grade), 'atualizado_siafe': _atualizado_siafe(grade),
            'totais': {'relatorio': total, 'detalhado': soma,
                       'confere': {k: _confere(soma[k], total[k]) for k in campos}}}


LEITORES = {'empenhos': ler_empenhos, 'restos': ler_restos}


# ---------------------------------------------------------------- gravação

def importar(tipo, caminho, arquivo, usuario):
    """Lê o relatório e substitui a importação anterior do mesmo tipo (numa
    transação). Retorna o resumo da leitura (linhas, totais, referência)."""
    lido = LEITORES[tipo](caminho)
    linhas = lido['linhas']
    s = schema()
    credores = {}
    for l in linhas:
        if l['credor_documento']:
            credores[l['credor_documento']] = l['credor_nome']
    with _cursor() as cur:
        if credores:
            execute_values(cur, f"""
                INSERT INTO {s}.credor (documento, nome) VALUES %s
                ON CONFLICT (documento) DO UPDATE SET nome = excluded.nome, atualizado_em = now()""",
                           sorted(credores.items()))
        cur.execute(f'DELETE FROM {s}.importacao_financeiro WHERE tipo = %s', [tipo])  # remove as linhas (cascade)
        cur.execute(f"""
            INSERT INTO {s}.importacao_financeiro
                (tipo, arquivo, referencia, atualizado_siafe, linhas, totais, importado_por)
            VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING id""",
                    [tipo, os.path.basename(arquivo), lido['referencia'], lido['atualizado_siafe'],
                     len(linhas), json.dumps(lido['totais'], ensure_ascii=False), usuario])
        imp = cur.fetchone()['id']
        if tipo == 'empenhos':
            execute_values(cur, f"""
                INSERT INTO {s}.empenho_aberto (importacao_id, empenho, fonte, credor_documento,
                                                valor_a_liquidar, valor_liquidado_a_pagar) VALUES %s""",
                           [(imp, l['empenho'], l['fonte'], l['credor_documento'] or None,
                             l['valor_a_liquidar'], l['valor_liquidado_a_pagar']) for l in linhas], page_size=1000)
        else:
            execute_values(cur, f"""
                INSERT INTO {s}.resto_pagar (importacao_id, empenho, processo, fonte, contrato, credor_documento,
                                             ordem_bancaria, data_ob, valor_inscrito, valor_pago, valor_a_pagar)
                VALUES %s""",
                           [(imp, l['empenho'], l['processo'], l['fonte'], l['contrato'], l['credor_documento'] or None,
                             l['ordem_bancaria'], l['data_ob'], l['valor_inscrito'], l['valor_pago'],
                             l['valor_a_pagar']) for l in linhas], page_size=1000)
    return lido


def importacoes():
    """{tipo: última importação (ou None)} para as telas."""
    resultado = dict.fromkeys(TIPOS)
    if not banco_painel.pronto():
        return resultado
    with _cursor() as cur:
        cur.execute(f"""SELECT tipo, arquivo, referencia, atualizado_siafe, linhas, totais, importado_em, importado_por
                        FROM {schema()}.importacao_financeiro""")
        for r in cur.fetchall():
            d = dict(r)
            d['importado_em'] = d['importado_em'].strftime('%d/%m/%Y %H:%M')
            resultado[d['tipo']] = d
    return resultado


# ---------------------------------------------------------------- página

def _f(v):
    return float(v) if v is not None else 0.0


def get_financeiro():
    """Dados da página: importações, totais, os dois relatórios cruzados por
    empenho, pagamentos por mês, valores por fonte e por credor."""
    imps = importacoes()
    vazio = {'importacoes': imps, 'empenhos': [], 'pagamentos_mes': [], 'fontes_restos': [],
             'fontes_empenhos': [], 'credores': [], 'kpis': {}}
    if not banco_painel.pronto() or not any(imps.values()):
        return vazio
    s = schema()
    with _cursor() as cur:
        # um registro por empenho, juntando os dois relatórios (pode estar em um só)
        cur.execute(f"""
            WITH e AS (
                SELECT empenho, MIN(fonte) AS fonte, MIN(credor_documento) AS documento,
                       SUM(valor_a_liquidar) AS a_liquidar, SUM(valor_liquidado_a_pagar) AS liquidado_a_pagar
                FROM {s}.empenho_aberto GROUP BY empenho
            ), r AS (
                SELECT empenho, MIN(processo) AS processo, MIN(fonte) AS fonte, MIN(contrato) AS contrato,
                       MIN(credor_documento) AS documento, SUM(valor_inscrito) AS inscrito,
                       SUM(valor_pago) AS pago, SUM(valor_a_pagar) AS a_pagar,
                       COUNT(ordem_bancaria) AS ordens, MAX(data_ob) AS ultima_ob
                FROM {s}.resto_pagar GROUP BY empenho
            )
            SELECT COALESCE(e.empenho, r.empenho) AS empenho,
                   COALESCE(e.documento, r.documento) AS documento, c.nome AS credor,
                   e.fonte AS fonte_empenho, r.fonte AS fonte_resto, r.processo, r.contrato,
                   e.a_liquidar, e.liquidado_a_pagar, r.inscrito, r.pago, r.a_pagar, r.ordens, r.ultima_ob,
                   e.empenho IS NOT NULL AS em_aberto, r.empenho IS NOT NULL AS em_restos
            FROM e FULL OUTER JOIN r ON r.empenho = e.empenho
            LEFT JOIN {s}.credor c ON c.documento = COALESCE(e.documento, r.documento)
            ORDER BY 1""")
        empenhos = []
        for x in cur.fetchall():
            d = dict(x)
            for k in ('a_liquidar', 'liquidado_a_pagar', 'inscrito', 'pago', 'a_pagar'):
                d[k] = round(_f(d[k]), 2) if d[k] is not None else None
            d['ultima_ob'] = d['ultima_ob'].strftime('%d/%m/%Y') if d['ultima_ob'] else None
            empenhos.append(d)

        cur.execute(f"""SELECT to_char(date_trunc('month', data_ob), 'YYYY-MM') AS mes,
                               SUM(valor_pago) AS pago, COUNT(*) AS ordens
                        FROM {s}.resto_pagar WHERE data_ob IS NOT NULL GROUP BY 1 ORDER BY 1""")
        pagamentos_mes = [{'mes': r['mes'], 'pago': round(_f(r['pago']), 2), 'ordens': r['ordens']}
                          for r in cur.fetchall()]

        cur.execute(f"""SELECT COALESCE(fonte, 'Sem fonte') AS fonte, SUM(valor_inscrito) AS inscrito,
                               SUM(valor_pago) AS pago, SUM(valor_a_pagar) AS a_pagar
                        FROM {s}.resto_pagar GROUP BY 1 ORDER BY 2 DESC""")
        fontes_restos = [{'fonte': r['fonte'], **{k: round(_f(r[k]), 2) for k in ('inscrito', 'pago', 'a_pagar')}}
                         for r in cur.fetchall()]

    # empenhos por fonte: totais do relatório (incluem os liquidados a pagar não detalhados)
    fontes_emp = (imps['empenhos'] or {}).get('totais', {}).get('fontes', {})
    fontes_empenhos = sorted(({'fonte': f, 'a_liquidar': round(v['valor_a_liquidar'], 2),
                               'liquidado_a_pagar': round(v['valor_liquidado_a_pagar'], 2)}
                              for f, v in fontes_emp.items() if v['valor_a_liquidar'] or v['valor_liquidado_a_pagar']),
                             key=lambda x: -(x['a_liquidar'] + x['liquidado_a_pagar']))

    # por credor: em aberto (a liquidar + liquidado a pagar) e restos a pagar
    por_credor = {}
    for e in empenhos:
        c = por_credor.setdefault(e['documento'] or '', {'documento': e['documento'], 'credor': e['credor'] or 'Sem credor',
                                                          'em_aberto': 0.0, 'inscrito': 0.0, 'pago': 0.0, 'a_pagar': 0.0,
                                                          'empenhos': 0})
        c['empenhos'] += 1
        c['em_aberto'] += (e['a_liquidar'] or 0) + (e['liquidado_a_pagar'] or 0)
        for k in ('inscrito', 'pago', 'a_pagar'):
            c[k] += e[k] or 0
    credores = sorted(({**c, **{k: round(c[k], 2) for k in ('em_aberto', 'inscrito', 'pago', 'a_pagar')}}
                       for c in por_credor.values()), key=lambda c: -(c['em_aberto'] + c['a_pagar']))

    tot_emp = (imps['empenhos'] or {}).get('totais', {}).get('relatorio', {})
    tot_rp = (imps['restos'] or {}).get('totais', {}).get('relatorio', {})
    inscrito = tot_rp.get('valor_inscrito') or 0
    kpis = {
        'a_liquidar': tot_emp.get('valor_a_liquidar'),
        'liquidado_a_pagar': tot_emp.get('valor_liquidado_a_pagar'),
        'empenhos_abertos': sum(e['em_aberto'] for e in empenhos),
        'rp_inscrito': tot_rp.get('valor_inscrito'),
        'rp_pago': tot_rp.get('valor_pago'),
        'rp_a_pagar': tot_rp.get('valor_a_pagar'),
        'rp_perc_pago': round(tot_rp.get('valor_pago', 0) / inscrito * 100, 1) if inscrito else None,
        'rp_empenhos': sum(e['em_restos'] for e in empenhos),
        'nos_dois': sum(e['em_aberto'] and e['em_restos'] for e in empenhos),
        'credores': len([c for c in credores if c['documento']]),
    }
    return {'importacoes': imps, 'empenhos': empenhos, 'pagamentos_mes': pagamentos_mes,
            'fontes_restos': fontes_restos, 'fontes_empenhos': fontes_empenhos, 'credores': credores,
            'kpis': kpis}


# ---------------------------------------------------------------- relatório (Excel)

SITUACOES = ('dois', 'aberto', 'restos', 'a_pagar')


def _so_letras_e_digitos(texto):
    """Texto para a busca: sem acento, minúsculo, sem ".", "-", "/" e espaços
    (o mesmo critério da busca da tabela na página)."""
    return re.sub(r'[.\-/\s]', '', _sem_acento(texto or ''))


def filtrar_empenhos(empenhos, busca='', situacao='', fonte=''):
    """Os filtros da tabela da página Financeiro, aplicados no servidor."""
    busca = _so_letras_e_digitos(busca)
    resultado = []
    for e in empenhos:
        if situacao == 'dois' and not (e['em_aberto'] and e['em_restos']):
            continue
        if situacao == 'aberto' and not (e['em_aberto'] and not e['em_restos']):
            continue
        if situacao == 'restos' and not (e['em_restos'] and not e['em_aberto']):
            continue
        if situacao == 'a_pagar' and not (e['a_pagar'] or 0) > 0:
            continue
        if fonte and fonte not in (e['fonte_empenho'], e['fonte_resto']):
            continue
        if busca:
            alvo = ' '.join(str(e.get(k) or '') for k in ('empenho', 'credor', 'documento', 'processo', 'contrato'))
            if busca not in _so_letras_e_digitos(alvo):
                continue
        resultado.append(e)
    return resultado


def _documento_formatado(d):
    if d and len(d) == 14:
        return f'{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}'
    if d and len(d) == 11:
        return f'{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}'
    return d or ''


# colunas da lista: (título, largura, campo ou função, formato)
_MOEDA = '#,##0.00'
_COLUNAS_RELATORIO = [
    ('Empenho', 16, 'empenho', None),
    ('Credor', 46, lambda e: e['credor'] or '', None),
    ('CNPJ / CPF', 20, lambda e: _documento_formatado(e['documento']), None),
    ('Fonte de recursos', 40, 'fonte_empenho', None),
    ('A liquidar', 16, lambda e: e['a_liquidar'] or 0, _MOEDA),
    ('Liquidado a pagar', 16, lambda e: e['liquidado_a_pagar'] or 0, _MOEDA),
    ('Total em aberto', 16, lambda e: (e['a_liquidar'] or 0) + (e['liquidado_a_pagar'] or 0), _MOEDA),
    ('Processo', 20, 'processo', None),
    ('Contrato', 14, 'contrato', None),
    ('Fonte (restos a pagar)', 14, 'fonte_resto', None),
    ('RP inscrito', 16, lambda e: e['inscrito'], _MOEDA),
    ('RP pago', 16, lambda e: e['pago'], _MOEDA),
    ('RP a pagar', 16, lambda e: e['a_pagar'], _MOEDA),
    ('Ordens bancárias', 10, 'ordens', '0'),
    ('Última OB', 12, 'ultima_ob', None),
]


def relatorio_empenhos_xlsx(busca='', situacao='', fonte=''):
    """Excel com a lista dos empenhos em aberto (último relatório importado),
    com os filtros da tabela da página e o cruzamento com os restos a pagar.
    Retorna (BytesIO, nome do arquivo)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    dados = get_financeiro()
    imp = dados['importacoes']['empenhos']
    lista = [e for e in filtrar_empenhos(dados['empenhos'], busca, situacao, fonte) if e['em_aberto']]
    lista.sort(key=lambda e: (-((e['a_liquidar'] or 0) + (e['liquidado_a_pagar'] or 0)), e['empenho']))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Empenhos em aberto'
    negrito = Font(bold=True)
    ws['A1'] = 'Relatório de Empenhos em Aberto'
    ws['A1'].font = Font(bold=True, size=14)
    info = [
        f'Período do relatório do SIAFE: {imp["referencia"] or "—"} (dados do SIAFE de {imp["atualizado_siafe"] or "—"})'
        if imp else 'Nenhum relatório de empenhos importado.',
        f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M")} pelo Painel de Indicadores.',
    ]
    filtros = [f'busca "{busca}"' if busca else '', f'fonte "{fonte}"' if fonte else '',
               {'dois': 'nos dois relatórios', 'aberto': 'só empenhos em aberto', 'restos': 'só restos a pagar',
                'a_pagar': 'com restos a pagar a pagar'}.get(situacao, '')]
    if any(filtros):
        info.append('Filtros: ' + '; '.join(f for f in filtros if f) + '.')
    if imp and not imp['totais']['confere'].get('valor_liquidado_a_pagar', True):
        total = imp['totais']['relatorio']['valor_liquidado_a_pagar']
        info.append('O relatório do SIAFE só detalha os empenhos com valor a liquidar: o liquidado a pagar total '
                    f'(R$ {total:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
                    + ') inclui empenhos que não aparecem nesta lista.')
    for i, texto in enumerate(info, start=2):
        ws.cell(row=i, column=1, value=texto)

    cab = len(info) + 3
    fundo = PatternFill('solid', start_color='007C6B')
    linha_fina = Side(style='thin', color='BFCFD3')
    for c, (titulo, largura, _, _) in enumerate(_COLUNAS_RELATORIO, start=1):
        cel = ws.cell(row=cab, column=c, value=titulo)
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = fundo
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = largura
    for r, e in enumerate(lista, start=cab + 1):
        for c, (_, _, campo, formato) in enumerate(_COLUNAS_RELATORIO, start=1):
            valor = campo(e) if callable(campo) else e.get(campo)
            cel = ws.cell(row=r, column=c, value=valor)
            cel.border = Border(bottom=linha_fina)
            if formato:
                cel.number_format = formato
    fim = cab + len(lista)
    total = fim + 1
    ws.cell(row=total, column=1, value=f'Total ({len(lista)} empenhos)').font = negrito
    for c, (_, _, _, formato) in enumerate(_COLUNAS_RELATORIO, start=1):
        if formato and lista:
            letra = get_column_letter(c)
            cel = ws.cell(row=total, column=c, value=f'=SUBTOTAL(9,{letra}{cab + 1}:{letra}{fim})')
            cel.number_format = formato
            cel.font = negrito
    ws.freeze_panes = ws.cell(row=cab + 1, column=2)
    if lista:
        ws.auto_filter.ref = f'A{cab}:{get_column_letter(len(_COLUNAS_RELATORIO))}{fim}'
    ws.row_dimensions[cab].height = 32

    arquivo = io.BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo, f'empenhos_em_aberto_{date.today().strftime("%Y-%m-%d")}.xlsx'
