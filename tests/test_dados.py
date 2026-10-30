"""Números do lançamento, API de custos, SIMAQ e configuração do banco."""
import pytest

import area_restrita
import banco_painel
import config_banco
import custos


# ---------------------------------------------------------------- num_br / ler_numero

@pytest.mark.parametrize('valor, formato, esperado', [
    (0.123456, 'decimal', '0,123456'),
    (1000, 'decimal', '1.000'),
    (2.5, 'inteiro', '2,5'),
    (3, 'inteiro', '3'),
    (1234.5, 'moeda', '1.234,50'),
    (1234.567, 'moeda', '1.234,567'),
    (None, 'decimal', ''),
    ('texto', 'texto', 'texto'),
])
def test_num_br(valor, formato, esperado):
    assert area_restrita.num_br(valor, formato) == esperado


@pytest.mark.parametrize('valor', [0.123456, 2.5, 1234.567, 13831, 0.651, 1000000.25])
def test_num_br_volta_igual_sem_arredondar(valor):
    """Reenviar o formulário sem mexer na linha não muda o valor gravado."""
    assert banco_painel.ler_numero(area_restrita.num_br(valor)) == valor


# ---------------------------------------------------------------- custos

LINHAS_CUSTOS = [
    {'mes': '2024-01', 'tipo': 'Pessoal', 'item': 'Salários', 'valor': 100.0},
    {'mes': '2024-02', 'tipo': 'Pessoal', 'item': 'Salários', 'valor': 110.0},
    {'mes': '2024-05', 'tipo': 'Insumos', 'item': 'Luvas', 'valor': 20.0},
]


@pytest.fixture
def dados_custos(monkeypatch):
    monkeypatch.setattr(banco_painel.CUSTOS, 'dados', lambda: [dict(l) for l in LINHAS_CUSTOS])


def test_custos_periodo(dados_custos):
    d = custos.get_custos('2024-01', '2024-05')
    assert d['meses_iso'] == ['2024-01', '2024-02', '2024-03', '2024-04', '2024-05']
    assert d['kpis']['total'] == 230.0
    assert d['kpis']['meses_com_dado'] == 3


def test_custos_padrao_ultimos_12_meses_completos(dados_custos):
    from datetime import date
    d = custos.get_custos(hoje=date(2024, 6, 15))
    assert (d['mes_ini'], d['mes_fim']) == ('2023-06', '2024-05')
    assert d['meses_disponiveis'][0] == '2023-06' and d['meses_disponiveis'][-1] == '2024-05'
    assert d['kpis']['total'] == 230.0


def test_custos_por_subcategoria(dados_custos, monkeypatch):
    monkeypatch.setattr(custos, '_subcategorias', lambda: {('Pessoal', 'Salários'): 'Servidores'})
    d = custos.get_custos('2024-01', '2024-05')
    assert d['por_subcategoria'] == [{'nome': 'Servidores', 'valor': 210.0},
                                     {'nome': 'Insumos (sem subcategoria)', 'valor': 20.0}]


@pytest.mark.parametrize('mes_ini, mes_fim', [(None, '2024-1'), ('abc', None), ('2024-13', None)])
def test_custos_mes_malformado(dados_custos, mes_ini, mes_fim):
    with pytest.raises(ValueError):
        custos.get_custos(mes_ini, mes_fim)


def test_api_custos_mes_malformado_da_400(cliente, dados_custos):
    r = cliente.get('/api/custos/dados?mes_fim=2024-1')
    assert r.status_code == 400 and not r.get_json()['ok']


def test_sem_banco_e_sem_planilha_fica_vazio(monkeypatch, tmp_path):
    monkeypatch.setattr(banco_painel, 'PASTA_DADOS', str(tmp_path))
    assert banco_painel.CUSTOS.dados() == []
    assert custos.get_custos() == {'vazio': True}


# ---------------------------------------------------------------- SIMAQ

def test_simaq_linha_so_com_resultado_vira_numerador():
    linha = {'numerador': None, 'denominador': None, 'resultado': 42.0, 'meta': None}
    banco_painel.SIMAQ._ajustar_linha_planilha(linha)
    assert linha['numerador'] == 42.0
    assert not banco_painel.SIMAQ.em_branco(linha)


class _CursorFalso:
    def __init__(self):
        self.fatores = {}

    def execute(self, sql, params):
        fator, _tipo, indicador = params
        self.fatores[indicador] = fator


def test_simaq_deduz_fator(monkeypatch, capsys):
    monkeypatch.setattr(banco_painel, 'schema', lambda: 'painel')

    def linha(indicador, n, d, r):
        return {'tipo': 'T', 'indicador': indicador, 'numerador': n, 'denominador': d, 'resultado': r}
    cur = _CursorFalso()
    banco_painel.SIMAQ._aplicar_fatores(cur, [
        linha('por mil', 2, 1000, 2.0), linha('por mil', 3, 1000, 3.0),
        linha('vezes 30', 1, 2, 15.0), linha('vezes 30', 2, 2, 30.0),
        linha('simples', 1, 2, 0.5),
        linha('irregular', 1, 2, 5.0), linha('irregular', 1, 2, 50.0),
    ])
    assert cur.fatores == {'por mil': 1000, 'vezes 30': 30.0}
    assert 'irregular' in capsys.readouterr().out  # vai para o log, para revisar em Cadastros


# ---------------------------------------------------------------- configuração do banco

def test_chave_malformada_nao_derruba(monkeypatch, tmp_path):
    arquivo = tmp_path / 'config_banco.enc'
    arquivo.write_bytes(b'conteudo-cifrado-qualquer')
    monkeypatch.setattr(config_banco, 'ARQUIVO_CONFIG', str(arquivo))
    monkeypatch.setattr(config_banco, 'ARQUIVO_LEGADO', str(tmp_path / 'nao-existe.json'))
    monkeypatch.setenv(config_banco.VARIAVEL_CHAVE, 'isto-nao-e-uma-chave-fernet')
    monkeypatch.setitem(config_banco._cache, 'lido_em', None)
    monkeypatch.setitem(config_banco._cache, 'mtime', None)

    cfg = config_banco.carregar()  # não levanta ValueError
    assert set(cfg) == set(config_banco.CAMPOS)
    assert config_banco.erro_arquivo()


def test_paginas_sem_dados_nao_quebram(cliente, tmp_path, monkeypatch):
    """Servidor novo: sem planilhas em dados/ e sem lançamentos."""
    import painel
    monkeypatch.setattr(banco_painel, 'PASTA_DADOS', str(tmp_path))
    monkeypatch.setattr(painel, 'PASTA_DADOS', str(tmp_path))
    assert cliente.get('/producao_hospitalar').status_code == 200
    assert cliente.get('/indicadores').status_code == 200
    assert cliente.get('/custos').status_code == 200
    assert cliente.get('/financeiro').status_code == 200
    r = cliente.get('/api/financeiro/dados').get_json()  # sem banco: página vazia, sem erro
    assert r['ok'] and r['dados']['empenhos'] == [] and r['dados']['importacoes'] == {'empenhos': None, 'restos': None}
    assert cliente.get('/faturamento').status_code == 404


def _planilha_siafe(caminho, linhas, mesclas=()):
    """xlsx com as linhas dadas (lista de listas, a partir da linha 1) e as
    mesclas ('A8:A10'...), no formato dos relatórios do SIAFE."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for linha in linhas:
        ws.append(linha)
    for m in mesclas:
        ws.merge_cells(m)
    wb.save(caminho)
    return str(caminho)


def test_financeiro_etl_credor_e_processo():
    import financeiro
    assert financeiro.separar_credor('54.858.014/0009-27 - ESSITY DO BRASIL - FILIAL') == \
        ('54858014000927', 'ESSITY DO BRASIL - FILIAL')  # só o primeiro " - " separa
    assert financeiro.separar_credor('SEM DOCUMENTO') == ('', 'SEM DOCUMENTO')
    assert financeiro.limpar_processo('NUP 24001.006317/2024-93 ') == '24001006317/202493'
    assert financeiro.limpar_processo(' - ') is None


def test_financeiro_le_empenhos_por_fonte(tmp_path):
    import financeiro
    cab = [['', 'Governo'], ['', 'Sefaz'], ['', 'Tesouro'],
           ['', 'Relatório de Empenhos em aberto por fonte de recursos - Período: Até 12 de 2025'], ['', ''],
           ['Despesa', '', 'Empenhos a Liquidar', 'Empenhos Liquidadas a Pagar'],
           ['241291 - HOSPITAL', '', 300.0, 70.0],
           ['   500 - RECURSOS NÃO VINCULADOS', '', 100.0, 50.0]]          # linha 8: fonte (mesclada A:B)
    dados = [['      2025NE000029', '54858014000927 - ESSITY', 100.0, 0.0],
             ['   600 - TRANSFERÊNCIAS', '', 200.0, 20.0],
             ['      2025NE000031', '20301535000100 - JB FARMA', 150.0, 20.0],
             ['      2025NE000032', '20301535000100 - JB FARMA', 50.0, 0.0],
             ['Impresso em 26/10/2026 - Dados atualizados em: 23/10/2026 10:41:05']]
    caminho = _planilha_siafe(tmp_path / 'empenhos.xlsx', cab + dados, ['A8:B8'])
    lido = financeiro.ler_empenhos(caminho)
    assert [(l['empenho'], l['fonte'][:3], l['credor_documento']) for l in lido['linhas']] == [
        ('2025NE000029', '500', '54858014000927'), ('2025NE000031', '600', '20301535000100'),
        ('2025NE000032', '600', '20301535000100')]
    assert lido['referencia'] == 'Até 12 de 2025' and lido['atualizado_siafe'] == '23/10/2026 10:41:05'
    t = lido['totais']
    assert t['confere'] == {'valor_a_liquidar': True, 'valor_liquidado_a_pagar': False}  # 50 não detalhado
    assert t['fontes']['500 - RECURSOS NÃO VINCULADOS']['valor_liquidado_a_pagar'] == 50.0


def test_financeiro_le_restos_repetindo_mescladas(tmp_path):
    import financeiro
    cab = [['', 'Governo'], ['', 'Sefaz'], ['', 'Tesouro'],
           ['', 'Restos a Pagar Inscritos, Pagos e A Pagar - UG: 241291 Até o Mês 10 / 2026'], [''],
           ['Unidade Gestora / Nota de Empenho', 'Unid. Orçam.', 'Processo', 'Fonte ', 'Contrato', 'Credor',
            'Ordem Bancária', 'Data Emissão OB', 'RP Processados', 'RP Não Processados', 'TOTAL INSCRITO',
            'VALOR PAGO', 'VALOR A PAGAR'],
           ['241291 - HOSPITAL', '', '', '', '', '', '', '', 0, 1000, 1000, 900, 100]]
    # um empenho com duas ordens bancárias: A a F mescladas nas 3 linhas (8 a 10)
    dados = [['   2024NE000272', '24200004214', 'NUP 24001.006317/2024-93', '6009200000', 'I001305848',
              '23596733000136 - SEROPEDICA', ' - ', ' - ', 0, 1000, 1000, 0, 100],
             ['', '', '', '', '', '', '2026OB002352', '24/03/2026', 0, 0, 0, 400, 0],
             ['', '', '', '', '', '', '2026OB002362', '25/03/2026', 0, 0, 0, 500, 0]]
    caminho = _planilha_siafe(tmp_path / 'restos.xlsx', cab + dados,
                              [f'{c}8:{c}10' for c in 'ABCDEF'])
    lido = financeiro.ler_restos(caminho)
    ls = lido['linhas']
    assert len(ls) == 3 and {l['empenho'] for l in ls} == {'2024NE000272'}
    assert {(l['processo'], l['fonte'], l['contrato'], l['credor_documento']) for l in ls} == \
        {('24001006317/202493', '6009200000', 'I001305848', '23596733000136')}  # repetidos nas mescladas
    assert [(l['ordem_bancaria'], str(l['data_ob'])) for l in ls] == [
        (None, 'None'), ('2026OB002352', '2026-03-24'), ('2026OB002362', '2026-03-25')]
    assert lido['referencia'] == 'Até o Mês 10 / 2026'
    assert all(lido['totais']['confere'].values())


def _empenho_cruzado(empenho, credor, doc, a_liquidar, em_restos=True, a_pagar=0.0, fonte='500 - RECURSOS'):
    return {'empenho': empenho, 'credor': credor, 'documento': doc, 'fonte_empenho': fonte,
            'fonte_resto': '5009100000' if em_restos else None, 'processo': '24001006317/202493' if em_restos else None,
            'contrato': None, 'a_liquidar': a_liquidar, 'liquidado_a_pagar': 0.0,
            'inscrito': a_liquidar if em_restos else None, 'pago': 0.0 if em_restos else None,
            'a_pagar': a_pagar if em_restos else None, 'ordens': 0, 'ultima_ob': None,
            'em_aberto': a_liquidar is not None, 'em_restos': em_restos}


def test_financeiro_filtros_do_relatorio():
    import financeiro
    es = [_empenho_cruzado('2025NE000001', 'ESSITY', '54858014000927', 100.0),
          _empenho_cruzado('2025NE000002', 'JOSÉ FARMÁCIA', '20301535000100', 50.0, em_restos=False, fonte='600 - SUS'),
          _empenho_cruzado('2024NE000009', 'ESSITY', '54858014000927', None, a_pagar=30.0)]
    nomes = lambda lista: [e['empenho'] for e in lista]  # noqa: E731
    assert nomes(financeiro.filtrar_empenhos(es, busca='jose farmacia')) == ['2025NE000002']  # sem acento
    assert nomes(financeiro.filtrar_empenhos(es, busca='54.858.014/0009-27')) == ['2025NE000001', '2024NE000009']
    assert nomes(financeiro.filtrar_empenhos(es, situacao='aberto')) == ['2025NE000002']
    assert nomes(financeiro.filtrar_empenhos(es, situacao='a_pagar')) == ['2024NE000009']
    assert nomes(financeiro.filtrar_empenhos(es, fonte='600 - SUS')) == ['2025NE000002']


def test_financeiro_relatorio_excel_so_empenhos_em_aberto(monkeypatch):
    import openpyxl
    import financeiro
    es = [_empenho_cruzado('2025NE000001', 'ESSITY', '54858014000927', 100.0),
          _empenho_cruzado('2025NE000002', 'JB FARMA', '20301535000100', 250.0, em_restos=False),
          _empenho_cruzado('2024NE000009', 'ESSITY', '54858014000927', None, a_pagar=30.0)]  # só restos: fora
    imp = {'referencia': 'Até 12 de 2025', 'atualizado_siafe': '23/10/2026 10:41:05',
           'totais': {'confere': {'valor_liquidado_a_pagar': False},
                      'relatorio': {'valor_liquidado_a_pagar': 2934525.93}}}
    monkeypatch.setattr(financeiro, 'get_financeiro',
                        lambda: {'empenhos': es, 'importacoes': {'empenhos': imp, 'restos': None}})
    arquivo, nome = financeiro.relatorio_empenhos_xlsx()
    assert nome.startswith('empenhos_em_aberto_') and nome.endswith('.xlsx')
    ws = openpyxl.load_workbook(arquivo).active
    textos = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    cab = textos.index('Empenho') + 1
    assert any('Até 12 de 2025' in (t or '') for t in textos[:cab])
    assert any('2.934.525,93' in (t or '') for t in textos[:cab])  # aviso dos liquidados não detalhados
    linhas = [[c.value for c in row] for row in ws.iter_rows(min_row=cab + 1, max_row=ws.max_row)]
    assert [l[0] for l in linhas[:-1]] == ['2025NE000002', '2025NE000001']  # do maior em aberto para o menor
    assert linhas[0][2] == '20.301.535/0001-00' and linhas[0][6] == 250.0
    assert linhas[-1][0] == 'Total (2 empenhos)' and linhas[-1][4].startswith('=SUBTOTAL(9,E')
    so_essity, _ = financeiro.relatorio_empenhos_xlsx(busca='essity')
    ws = openpyxl.load_workbook(so_essity).active
    assert sum(1 for row in ws.iter_rows(values_only=True) if row[0] and str(row[0]).startswith('20')) == 1


def test_financeiro_recusa_outra_planilha(tmp_path):
    import financeiro
    import pytest
    caminho = _planilha_siafe(tmp_path / 'outra.xlsx', [['a', 'b']] * 8)
    with pytest.raises(financeiro.PlanilhaInvalida):
        financeiro.ler_restos(caminho)
    with pytest.raises(financeiro.PlanilhaInvalida):
        financeiro.ler_empenhos(caminho)


def test_subcategoria_cooperativa_pelos_nomes_dos_itens():
    """Só os itens citados entram em Cooperativa; os demais ficam sem subcategoria."""
    nomes = ['CEMERG', 'CCÁRDIO', 'COAPH (Enfermeiros)', 'COAPH (Médicos Generalistas)', 'COMINT',
             'COOPANEST', 'COOPCARDIO', 'COOPCLÍNIC', 'HEMOCOOP', 'Cooperativas', 'Salário / SESA']
    dentro = [n for n in nomes if banco_painel._chave_nome(n).startswith(banco_painel.SUBCATEGORIA_COOPERATIVA)]
    assert dentro == nomes[:6]


def test_cadastro_item_custo_tem_subcategoria_opcional():
    item = banco_painel.CADASTROS_POR_TABELA['item_custo']
    assert item.pais == [('tipo_custo', 'Categoria')]
    assert item.opcionais == [('subcategoria_custo', 'Subcategoria', 'Sem subcategoria', 'tipo_custo')]
    assert banco_painel.CADASTROS_POR_TABELA['tipo_custo'].titulo == 'Categorias de Custo'


def test_classificacao_padrao_sem_arquivo_tenta_de_novo(monkeypatch):
    # sem o CSV, a carga não é registrada em migracao (roda no próximo início)
    class Cursor:
        def __init__(self):
            self.sqls = []

        def execute(self, sql, params=None):
            self.sqls.append(sql)

        def fetchone(self):
            return None
    cur = Cursor()
    monkeypatch.setattr(banco_painel, '_ler_classificacoes_padrao', lambda: [])
    banco_painel._carregar_classificacoes_padrao(cur)
    assert not any('INSERT' in s for s in cur.sqls)
    # o arquivo do sistema vai junto com o projeto (pasta sql/)
    monkeypatch.undo()
    assert len(banco_painel._ler_classificacoes_padrao()) > 900


def test_planilha_mes_em_texto_iso_e_brasileiro(tmp_path):
    import pandas as pd
    caminho = tmp_path / 'producao.xlsx'
    pd.DataFrame([['2026-08-01', 'IMAGEM', 'EXAME', 'Tomografia', 10],   # ISO: agosto (não 8 de janeiro)
                  ['01/09/2026', 'IMAGEM', 'EXAME', 'Tomografia', 20],   # dia primeiro: setembro
                  ['2026-10', 'IMAGEM', 'EXAME', 'Tomografia', 30]],
                 columns=['Mês', 'Serviço', 'Tipo', 'Procedimento', 'Qtd']).to_excel(caminho, index=False)
    linhas = banco_painel.PRODUCAO.ler_planilha(str(caminho))
    assert [(l['mes'], l['quantidade']) for l in linhas] == [('2026-08', 10), ('2026-09', 20), ('2026-10', 30)]


def test_importacao_cadastra_itens_novos(monkeypatch):
    t = banco_painel.SIMAQ
    linhas = [{'mes': '2026-09', 'tipo': 'ESTRATÉGICO', 'indicador': 'Novo', 'numerador': 1, 'denominador': 2},
              {'mes': '2026-09', 'tipo': 'ESTRATÉGICO', 'indicador': 'Antigo', 'numerador': 3, 'denominador': 4}]
    chamadas = {}
    monkeypatch.setattr(t, 'ler_planilha', lambda caminho: linhas)
    monkeypatch.setattr(t, 'itens_faltando', lambda ls: [('ESTRATÉGICO', 'Novo')])
    monkeypatch.setattr(t, 'salvar', lambda ls, usuario, criar_cadastros: chamadas.update(criar=criar_cadastros) or len(ls))
    monkeypatch.setattr(t, '_depois_de_cadastrar', lambda ls: chamadas.update(fator=[l['indicador'] for l in ls]))
    assert t.importar_cadastrando('x.xlsx', 'teste') == (2, [('ESTRATÉGICO', 'Novo')])
    assert chamadas == {'criar': True, 'fator': ['Novo']}  # fator deduzido só do indicador novo
    import area_restrita
    novos = [('ESTRATÉGICO', f'I{n}') for n in range(7)]
    assert area_restrita._texto_itens_novos([]) == ''
    assert '7 item(ns) novo(s)' in area_restrita._texto_itens_novos(novos) and 'e mais 2' in area_restrita._texto_itens_novos(novos)


def test_cadastro_tipo_consulta_com_codigo_do_sistema():
    cad = banco_painel.CADASTROS_POR_TABELA['tipo_consulta']
    assert cad.titulo == 'Tipos de Consulta' and cad.grupo == 'Ambulatório'
    assert [(c.nome, c.formato) for c in cad.extras] == [('isn_tipo_consulta', 'inteiro')]
    # carga inicial: arquivo do sistema em sql/ (um nome por código, sem repetir)
    tipos = banco_painel.tipos_consulta_padrao()
    assert tipos[34] == 'RETORNO' and tipos[31] == '1ª consulta' and tipos[29] == 'EXTRA'
    assert tipos[36] == 'TELEATENDIMENTO' and tipos[38] == 'REGULAÇÃO'
    assert len(set(tipos.values())) == len(tipos) == 12


def test_tipos_consulta_sem_arquivo_tenta_de_novo(monkeypatch):
    class Cursor:
        def __init__(self):
            self.sqls = []

        def execute(self, sql, params=None):
            self.sqls.append(sql)

        def fetchone(self):
            return None
    cur = Cursor()
    monkeypatch.setattr(banco_painel, 'tipos_consulta_padrao', lambda: {})
    banco_painel._criar_tipos_consulta(cur)
    assert not any('INSERT' in s for s in cur.sqls)  # nem a carga nem o registro em migracao


def test_lancamento_custos_agrupa_por_categoria_e_subcategoria():
    import lancamento_custos
    linhas = [
        {'tipo': 'Pessoal', 'subcategoria': 'Cooperativa', 'item': 'CEMERG', 'valor': 10.0},
        {'tipo': 'Pessoal', 'subcategoria': 'Cooperativa', 'item': 'COMINT', 'valor': 5.0},
        {'tipo': 'Pessoal', 'subcategoria': None, 'item': 'RESIDENTES', 'valor': None},
        {'tipo': 'Despesas Gerais', 'subcategoria': None, 'item': 'Hemodiálise', 'valor': 2.0},
    ]
    pessoal, gerais = lancamento_custos.agrupar(linhas, 'valor')
    assert pessoal['com_sub'] and pessoal['total'] == 15
    assert [(s['nome'], s['total'], len(s['itens'])) for s in pessoal['subs']] == [
        ('Cooperativa', 15, 2), ('Sem subcategoria', 0, 1)]
    assert not gerais['com_sub'] and gerais['total'] == 2


def test_ambulatorio_consultas_por_status_por_sala(monkeypatch):
    import ambulatorio
    linhas = [
        {'sala': 'SL-01', 'especialidade': 'Cardiologia', 'status': 'CONCLUÍDO',
         'sexo': 'Feminino', 'municipio': 'FORTALEZA/CE', 'qtd': 5},
        {'sala': 'SL-01', 'especialidade': 'Cardiologia', 'status': 'FALTOU',
         'sexo': 'Feminino', 'municipio': 'FORTALEZA/CE', 'qtd': 2},
        {'sala': 'SL-02', 'especialidade': 'Pneumologia', 'status': 'CONCLUÍDO',
         'sexo': 'Masculino', 'municipio': 'SOBRAL/CE', 'qtd': 9},
    ]
    monkeypatch.setattr(ambulatorio, '_salas_periodo', lambda ini, fim: linhas)
    d = ambulatorio.get_salas('2026-10-01', '2026-10-31')
    assert d['salas'] == ['SL-02', 'SL-01']  # da mais para a menos usada
    assert d['series'] == [{'status': 'CONCLUÍDO', 'valores': [9, 5]}, {'status': 'FALTOU', 'valores': [0, 2]}]
    d = ambulatorio.get_salas('2026-10-01', '2026-10-31', especialidade='Cardiologia')
    assert d['salas'] == ['SL-01']
    d = ambulatorio.get_salas('2026-10-01', '2026-10-31', sexo='Masculino')
    assert d['salas'] == ['SL-02'] and d['series'] == [{'status': 'CONCLUÍDO', 'valores': [9]}]
    d = ambulatorio.get_salas('2026-10-01', '2026-10-31', municipio='FORTALEZA/CE')
    assert d['salas'] == ['SL-01']


def test_ambulatorio_mensal(monkeypatch):
    from datetime import date
    import ambulatorio
    s, o = date(2026, 9, 10), date(2026, 10, 5)
    visitas = [  # (paciente, especialidade, status, dia, município, sexo, idade, internação, tipo de consulta)
        (1, 'Cardiologia', 'CONCLUÍDO', s, 'FORTALEZA/CE', 'Feminino', 50, 11, 31),
        (1, 'Cardiologia', 'FALTOU', date(2026, 9, 11), 'FORTALEZA/CE', 'Feminino', 50, 12, 34),
        (2, 'Pneumologia', 'ALTA AMBULATORIAL', o, 'SOBRAL/CE', 'Masculino', 30, 13, 34),
        (3, 'Cardiologia', 'CONCLUÍDO', o, 'SOBRAL/CE', 'Feminino', 40, 14, 34),
    ]
    consultas = [  # (especialidade, status, hora, minutos, sexo, município, dia)
        ('Cardiologia', 'CONCLUÍDO', 8, 10.0, 'Feminino', 'FORTALEZA/CE', s),
        ('Pneumologia', 'ALTA AMBULATORIAL', 9, 20.0, 'Masculino', 'SOBRAL/CE', o),
        ('Cardiologia', 'CONCLUÍDO', 9, 30.0, 'Feminino', 'SOBRAL/CE', o),
    ]
    monkeypatch.setattr(ambulatorio, '_visitas_periodo', lambda ini, fim: visitas)
    monkeypatch.setattr(ambulatorio, '_consultas_periodo', lambda ini, fim: consultas)
    monkeypatch.setattr(ambulatorio.banco_painel, 'macrorregiao_dos_municipios',
                        lambda: {'FORTALEZA': 'Fortaleza', 'SOBRAL': 'Norte'})
    monkeypatch.setattr(ambulatorio.banco_painel, 'nomes_tipos_consulta', lambda: {34: 'RETORNO'})
    d = ambulatorio.get_mensal('2026-10', '2026-09')  # meses em qualquer ordem
    # por tipo de consulta: mês a mês, a tabela do período e as linhas do comparativo
    assert d['tipo_consulta'] == [{'nome': 'RETORNO', 'valores': [1, 2]}, {'nome': 'Código 31', 'valores': [1, 0]}]
    assert d['tipo_especialidade']['tipos'] == ['RETORNO', 'Código 31']
    assert d['tipo_especialidade']['linhas'][0] == {'nome': 'Cardiologia', 'valores': [2, 1], 'total': 3}
    grupo = [c for c in d['_campos'] if len(c) == 5]
    assert [c[4] for c in grupo] == [False, True, True, False, True, True, False, True]  # tipo: detalhe
    assert [c[0] for c in grupo] == [
        'tipoesp_total', 'tipoesp_total_tipo_RETORNO', 'tipoesp_total_tipo_Código 31',
        'tipoesp_esp_Cardiologia', 'tipoesp_esp_Cardiologia_tipo_RETORNO', 'tipoesp_esp_Cardiologia_tipo_Código 31',
        'tipoesp_esp_Pneumologia', 'tipoesp_esp_Pneumologia_tipo_RETORNO']
    assert d['tipoesp_esp_Cardiologia'] == [2, 1] and d['tipoesp_esp_Pneumologia_tipo_RETORNO'] == [0, 1]
    assert d['meses_iso'] == ['2026-09', '2026-10']
    assert d['atendimentos'] == [2, 2] and d['prontuarios'] == [1, 2]
    assert d['media_diaria'] == [1.0, 2.0]  # 2 dias com atendimento em set, 1 em out
    assert d['concluidos'] == [1, 2] and d['taxa_absenteismo'] == [50.0, 0.0] and d['taxa_alta'] == [0.0, 50.0]
    assert d['consultas'] == [1, 2] and d['tempo_medio'] == [10.0, 25.0]
    assert d['sexo'] == [{'nome': 'Feminino', 'valores': [2, 1]}, {'nome': 'Masculino', 'valores': [0, 1]}]
    assert d['atendimentos_status_FALTOU'] == [1, 0]
    assert d['_campos'][0][0] == 'atendimentos' and d['_campos'][1][0].startswith('atendimentos_status_')
    # filtro de macrorregião: só Sobral (Norte), em outubro
    n = ambulatorio.get_mensal('2026-09', '2026-10', macrorregiao='Norte')
    assert n['atendimentos'] == [0, 2] and n['consultas'] == [0, 2]
    assert n['media_diaria'] == [0.0, 2.0]  # a média continua pelos dias em que o ambulatório atendeu
    assert ambulatorio.get_mensal('2026-09', '2026-10', especialidade='Pneumologia')['atendimentos'] == [0, 1]


def test_ambulatorio_por_tipo_especialidade():
    from datetime import date
    import ambulatorio
    d = date(2026, 10, 1)
    visitas = [  # (paciente, especialidade, status, dia, município, sexo, idade, internação, tipo)
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino', 50, 11, 34),
        (2, 'Cardiologia', 'FALTOU', d, 'FORTALEZA/CE', 'Feminino', 50, 12, 31),
        (3, 'Pneumologia', 'CONCLUÍDO', d, 'SOBRAL/CE', 'Masculino', 30, 13, 34),
        (4, 'Pneumologia', 'CONCLUÍDO', d, 'SOBRAL/CE', 'Masculino', 30, 14, 34),
        (5, 'Pneumologia', 'CONCLUÍDO', d, 'SOBRAL/CE', 'Masculino', 30, 15, None),
    ]
    nomes = {34: 'RETORNO', 31: '1ª consulta'}
    t = ambulatorio.por_tipo_especialidade(visitas, nomes)
    assert t['tipos'] == ['RETORNO', '1ª consulta', 'Sem tipo']
    assert t['linhas'] == [{'nome': 'Pneumologia', 'valores': [2, 0, 1], 'total': 3},
                           {'nome': 'Cardiologia', 'valores': [1, 1, 0], 'total': 2}]
    assert t['total'] == {'nome': 'Total', 'valores': [3, 1, 1], 'total': 5}
    assert ambulatorio.por_tipo_especialidade(visitas, nomes, status='FALTOU')['tipos'] == ['1ª consulta']
    assert ambulatorio.nome_tipo_consulta(99, nomes) == 'Código 99'


def test_comparativo_individual_linha_so_no_mes_b_fica_no_lugar(monkeypatch):
    import painel
    campos_a = [('a', 'A', 'qtd'), ('a_status_X', 'X', 'qtd'), ('b', 'B', 'qtd')]
    campos_b = [('a', 'A', 'qtd'), ('a_status_Y', 'Y', 'qtd'), ('b', 'B', 'qtd')]

    def mensal(ini, fim):
        um = ini == '2026-09'
        return {'meses': [ini], 'meses_iso': [ini], 'a': [1], 'b': [2],
                'a_status_X' if um else 'a_status_Y': [1], '_campos': campos_a if um else campos_b}
    with painel.app.test_request_context('/?mes_a=2026-09&mes_b=2026-10&modo=individual'):
        r = painel._hcasg_comparativo_resposta(mensal, campos_a).get_json()
    assert [l['chave'] for l in r['linhas']] == ['a', 'a_status_Y', 'a_status_X', 'b']
    assert all(l['grupo'] is None for l in r['linhas'])


def test_comparativo_detalhe_explicito_no_campo():
    import painel
    dados = {'meses': ['Out/26'], 'a_status_X': [1], 'esp_A_status_B': [2], 'esp_A_status_B_x': [3]}
    campos = [('a_status_X', 'X', 'qtd'),                         # 3 elementos: detalhe pela chave
              ('esp_A_status_B', 'A_status_B', 'qtd', 'g', False),  # explícito: o nome não decide
              ('esp_A_status_B_x', 'x', 'qtd', 'g', True)]
    linhas = painel._hcasg_montar_comparativo_periodo(dados, campos)['linhas']
    assert [(l['grupo'], l['sub']) for l in linhas] == [(None, True), ('g', False), ('g', True)]


def test_ambulatorio_pedidos_simultaneos_fazem_uma_consulta():
    import threading
    import time
    import ambulatorio
    chamadas = []

    def calcular():
        chamadas.append(1)
        time.sleep(0.2)  # consulta demorada no banco
        return ['ok']
    chave = ('teste-simultaneo', time.monotonic())
    resultados = []
    fios = [threading.Thread(target=lambda: resultados.append(ambulatorio._em_cache(chave, calcular)))
            for _ in range(5)]
    for f in fios:
        f.start()
    for f in fios:
        f.join()
    assert chamadas == [1] and resultados == [['ok']] * 5


def test_ambulatorio_filtro_sexo_municipio_macrorregiao():
    import ambulatorio
    mapa = {'FORTALEZA': '1ª Fortaleza', 'SOBRAL': '2ª Sobral'}
    assert ambulatorio.filtro_paciente() is None
    f = ambulatorio.filtro_paciente(sexo='Feminino')
    assert f('Feminino', 'SOBRAL/CE') and not f('Masculino', 'SOBRAL/CE')
    f = ambulatorio.filtro_paciente(macrorregiao='2ª Sobral', macro_do_municipio=mapa)
    assert f('Masculino', 'SOBRAL/CE') and not f('Masculino', 'FORTALEZA/CE')
    f = ambulatorio.filtro_paciente(macrorregiao='Outros estados', macro_do_municipio=mapa)
    assert f('Feminino', 'SOBRAL/PI') and not f('Feminino', 'SOBRAL/CE')
    f = ambulatorio.filtro_paciente(sexo='Feminino', municipio='FORTALEZA/CE')
    assert f('Feminino', 'FORTALEZA/CE') and not f('Feminino', 'SOBRAL/CE') and not f('Masculino', 'FORTALEZA/CE')
    assert ambulatorio.macrorregiao_do_municipio('Sem município', mapa) == 'Sem município'
    assert ambulatorio.macrorregiao_do_municipio('IGUATU/CE', mapa) == 'Ceará (não classificado)'
    # atendimentos antigos: município sem a UF
    assert ambulatorio.macrorregiao_do_municipio('SOBRAL', mapa) == '2ª Sobral'
    assert ambulatorio.macrorregiao_do_municipio('XYZ', mapa) == 'UF não informada'


def test_ambulatorio_por_prontuario_e_intervalo_medio():
    from datetime import date
    import ambulatorio
    visitas = [
        (1, 'Cardiologia', 'CONCLUÍDO', date(2026, 1, 10), 'FORTALEZA/CE'),
        (1, 'Cardiologia', 'CONCLUÍDO', date(2026, 4, 10), 'FORTALEZA/CE'),   # 90 dias depois
        (1, 'Cardiologia', 'FALTOU', date(2026, 4, 10), 'FORTALEZA/CE'),      # mesmo dia: não é retorno
        (2, 'Cardiologia', 'CONCLUÍDO', date(2026, 2, 1), 'CAUCAIA/CE'),
        (1, 'Pneumologia', 'CONCLUÍDO', date(2026, 3, 1), 'FORTALEZA/CE'),
    ]
    linhas, total = ambulatorio.por_prontuario(visitas)
    cardio = linhas[0]
    assert (cardio['nome'], cardio['atendimentos'], cardio['prontuarios'], cardio['com_retorno']) == ('Cardiologia', 4, 2, 1)
    assert cardio['por_prontuario'] == 2.0 and cardio['intervalo_meses'] == round(90 / 30.44, 1)
    assert linhas[1]['intervalo_meses'] is None  # sem retorno
    assert (total['atendimentos'], total['prontuarios']) == (5, 2)  # paciente 1 conta uma vez no total
    assert ambulatorio.por_municipio(visitas) == [
        {'nome': 'FORTALEZA/CE', 'qtd': 4, 'prontuarios': 1}, {'nome': 'CAUCAIA/CE', 'qtd': 1, 'prontuarios': 1}]
    assert ambulatorio.por_municipio(visitas, especialidade='Pneumologia') == [
        {'nome': 'FORTALEZA/CE', 'qtd': 1, 'prontuarios': 1}]
    linhas, _ = ambulatorio.por_prontuario(visitas, status='FALTOU')
    assert [(l['nome'], l['atendimentos']) for l in linhas] == [('Cardiologia', 1)]


def test_ambulatorio_por_macrorregiao():
    from datetime import date
    import ambulatorio
    d = date(2026, 10, 1)
    visitas = [
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE'),
        (2, 'Cardiologia', 'CONCLUÍDO', d, 'SOBRAL/CE'),
        (3, 'Cardiologia', 'CONCLUÍDO', d, 'FLORES/CE'),      # do Ceará, fora do cadastro
        (4, 'Cardiologia', 'CONCLUÍDO', d, 'MOSSORO/RN'),
        (5, 'Cardiologia', 'CONCLUÍDO', d, 'Sem município'),
    ]
    mapa = {'FORTALEZA': 'Fortaleza', 'SOBRAL': 'Norte'}
    assert {m['nome']: m['qtd'] for m in ambulatorio.por_macrorregiao(visitas, mapa)} == {
        'Fortaleza': 1, 'Norte': 1, 'Ceará (não classificado)': 1, 'Outros estados': 1, 'Sem município': 1}


def test_ambulatorio_por_sexo_e_especialidade():
    from datetime import date
    import ambulatorio
    d = date(2026, 10, 1)
    visitas = [
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino'),
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino'),  # mesmo prontuário
        (2, 'Cardiologia', 'FALTOU', d, 'FORTALEZA/CE', 'Masculino'),
        (3, 'Pneumologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Não informado'),
        (4, 'Pneumologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Masculino'),
    ]
    s = ambulatorio.por_sexo(visitas)
    assert s['sexos'] == ['Feminino', 'Masculino', 'Não informado']
    assert [(x['nome'], x['qtd'], x['prontuarios']) for x in s['total']] == [
        ('Feminino', 2, 1), ('Masculino', 2, 2), ('Não informado', 1, 1)]
    assert s['por_especialidade'] == [
        {'nome': 'Cardiologia', 'qtd': 3, 'sexos': [2, 1, 0]},
        {'nome': 'Pneumologia', 'qtd': 2, 'sexos': [0, 1, 1]}]
    assert ambulatorio.por_sexo(visitas, status='FALTOU')['sexos'] == ['Masculino']


def test_ambulatorio_por_idade_um_ano_por_barra():
    from datetime import date
    import ambulatorio
    d = date(2026, 10, 1)
    visitas = [
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino', 2),
        (1, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino', 2),   # mesmo prontuário
        (2, 'Cardiologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Masculino', 0),
        (3, 'Pneumologia', 'FALTOU', d, 'FORTALEZA/CE', 'Masculino', 5),
        (4, 'Pneumologia', 'CONCLUÍDO', d, 'FORTALEZA/CE', 'Feminino', None),  # sem data de nascimento
    ]
    r = ambulatorio.por_idade(visitas)
    assert [(x['nome'], x['qtd'], x['prontuarios']) for x in r['idades']] == [
        ('0', 1, 1), ('1', 0, 0), ('2', 2, 1), ('3', 0, 0), ('4', 0, 0), ('5', 1, 1), ('Não informada', 1, 1)]
    assert r['media'] == round((2 + 2 + 0 + 5) / 4, 1) and r['mediana'] == 2
    assert [x['nome'] for x in ambulatorio.por_idade(visitas, status='FALTOU')['idades']][-1] == '5'
    assert ambulatorio.por_idade([]) == {'idades': [], 'media': None, 'mediana': None}


def test_ambulatorio_tempo_medio_por_consulta():
    import ambulatorio
    consultas = [
        ('Cardiologia', 'CONCLUÍDO', 7, 10.0),
        ('Cardiologia', 'CONCLUÍDO', 7, 20.0),
        ('Cardiologia', 'ALTA AMBULATORIAL', 8, 60.0),
        ('Pneumologia', 'CONCLUÍDO', 9, 5.0),
        ('Pneumologia', 'CONCLUÍDO', 9, None),   # sem fim: fora do tempo
    ]
    t = ambulatorio.tempo_consulta(consultas)
    assert t['especialidades'] == [
        {'nome': 'Cardiologia', 'consultas': 3, 'media': 30.0, 'mediana': 20.0},
        {'nome': 'Pneumologia', 'consultas': 1, 'media': 5.0, 'mediana': 5.0}]
    assert t['total'] == {'nome': 'Total', 'consultas': 4, 'media': 23.8, 'mediana': 15.0}
    assert ambulatorio.tempo_consulta(consultas, status='ALTA AMBULATORIAL')['total']['media'] == 60.0
    assert ambulatorio.tempo_consulta([]) == {'especialidades': [], 'total': None}


def test_macrorregioes_do_ceara_tem_os_184_municipios():
    m = banco_painel.MACRORREGIOES_CE
    assert list(m) == ['Fortaleza', 'Norte', 'Cariri', 'Sertão Central', 'Litoral Leste/Jaguaribe']
    nomes = [n for v in m.values() for n in v]
    assert len(nomes) == len(set(nomes)) == 184


def test_classificacao_padrao_dos_exames_vem_com_o_sistema():
    linhas = banco_painel._ler_classificacoes_padrao()
    assert len(linhas) >= 900
    codigos = [i for i, _, _ in linhas]
    assert len(codigos) == len(set(codigos))  # um exame, uma classificação
    assert {c for _, _, c in linhas} <= set(banco_painel.CLASSIFICACOES_EXAME_PADRAO)
    # o mesmo nome nunca com classificações diferentes (a carga também casa pelo nome)
    por_nome = {}
    for _, nome, c in linhas:
        assert nome and por_nome.setdefault(nome.upper(), c) == c


def test_macrorregioes_sem_banco_pronto_usa_a_lista_oficial(monkeypatch):
    # logo depois de iniciar o servidor o banco do painel ainda está conectando:
    # os municípios continuam classificados
    monkeypatch.setattr(banco_painel, 'pronto', lambda: False)
    mapa = banco_painel.macrorregiao_dos_municipios()
    assert len(mapa) == 184
    assert mapa['FORTALEZA'] == 'Fortaleza' and mapa['SOBRAL'] == 'Norte' and mapa['VICOSA DO CEARA'] == 'Norte'
