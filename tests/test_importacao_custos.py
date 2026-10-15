"""Leitura da aba GERAL da planilha do setor de custos (sem banco)."""
import io
import os

import openpyxl
import pytest

import importacao_custos as ic

PLANILHA_REAL = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'dados', 'Planilha-do-mes-de-Junho-2026.xlsx')


def _planilha(tmp_path, linhas, aba='GERAL '):
    """Planilha no formato do setor: grupos mesclados na linha 9, cabeçalho na
    10 (A..D + unidades da E em diante) e as linhas dadas a partir da 11."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    ws['E7'] = 'MÊS: JUNHO 2026'
    ws['E9'] = 'INTERNAÇÕES'
    ws.merge_cells('E9:F9')
    ws['G9'] = 'AUXILIARES'
    for col, v in zip('ABCDEFG', ['1. CUSTOS DIRETOS', 'Item', 'Geral', 'valor', 'Unidade A', 'UTI\nAdulto', 'Farmácia']):
        ws[f'{col}10'] = v
    for n, linha in enumerate(linhas, start=11):
        for col, v in zip('ABCDEFG', linha):
            ws[f'{col}{n}'] = v
    caminho = tmp_path / 'planilha.xlsx'
    wb.save(caminho)
    return str(caminho)


LINHAS = [
    ['1.1 - PESSOAL', '%', '%', 'Valor (R$)', 'Valor (R$)', 'Valor (R$)', 'Valor (R$)'],
    ['1.1.1 – Salário / SESA      ', 0.5, 0.5, 300, 100, 150, 50],
    ['1.1.7.1 – COMINT', 0.5, 0.5, 100, None, 100, None],
    ['1.1.4 – SUB-TOTAL', 1, 1, 400, 100, 250, 50],
    ['1.2   –  MATERIAL DE CONSUMO'],
    ['1.2.1 – Mat. médico hospitalar', 1, 1, 60, 20, 40, None],
    ['SUB-TOTAL', 1, 1, 60, 20, 40, None],
    ['2 - CUSTOS INDIRETOS'],
    ['2.2 – Energia elétrica', 1, 1, 40, 10, 10, 20],
    ['5 - CUSTO TOTAL', 1, 1, 500, 130, 300, 70],
    ['6 Rateio dos Custos  Aux. e Adm.'],
    ['6.7 – Farmácia', 1, 1, 70, 30, 40, 70],
    ['6.12 – Total Rateios Centros Aux Adm.', 1, 1, 70, 30, 40, 0],
    ['7. - CUSTO TOTAL C/ RATEIOS', None, None, None, 160, 340, 0],
    ['8. -  PRODUÇÃO', None, None, None, 10, 20, None],
    ['9. - CUSTO UNITÁRIO FINAL', None, None, None, 16, 17, None],
    [None, None, None, None, 'Diária', 'Diária', None],
]


def test_le_itens_unidades_rateio_e_producao(tmp_path):
    d = ic.ler_geral(_planilha(tmp_path, LINHAS))
    assert d['mes_planilha'] == 'JUNHO 2026'
    assert d['custo_total_planilha'] == 500
    assert [(u['nome'], u['grupo'], u['medida'], u['producao']) for u in d['unidades']] == [
        ('Unidade A', 'INTERNAÇÕES', 'Diária', 10), ('UTI Adulto', 'INTERNAÇÕES', 'Diária', 20),
        ('Farmácia', 'AUXILIARES', None, None)]
    custos = [(i['tipo'], i['item'], i['total']) for i in d['itens'] if i['categoria'] == 'custo']
    assert custos == [('Pessoal', 'Salário / SESA', 300), ('Pessoal', 'COMINT', 100),
                      ('Material de Consumo', 'Mat. médico hospitalar', 60),
                      ('Custos Indiretos', 'Energia elétrica', 40)]  # subtotais e totais de fora
    rateio = [i for i in d['itens'] if i['categoria'] == 'rateio']
    assert [(i['item'], i['total'], i['valores']) for i in rateio] == [
        ('Farmácia', 70, {'Unidade A': 30, 'UTI Adulto': 40, 'Farmácia': 70})]
    assert d['itens'][1]['valores'] == {'UTI Adulto': 100}  # célula vazia não vira valor
    assert d['avisos'] == []


def test_avisa_total_diferente_da_soma_das_unidades(tmp_path):
    linhas = [l[:] for l in LINHAS]
    linhas[1][3] = 999  # coluna D não bate com E+F+G
    d = ic.ler_geral(_planilha(tmp_path, linhas))
    assert any('Salário / SESA' in a and 'soma das unidades' in a for a in d['avisos'])
    assert any('CUSTO TOTAL' in a for a in d['avisos'])


def test_sem_aba_geral(tmp_path):
    with pytest.raises(ic.ErroPlanilha, match='aba GERAL'):
        ic.ler_geral(_planilha(tmp_path, LINHAS, aba='Plan1'))


def test_sem_cabecalho(tmp_path):
    caminho = _planilha(tmp_path, LINHAS)
    wb = openpyxl.load_workbook(caminho)
    wb.active['D10'] = 'outra coisa'
    wb.save(caminho)
    with pytest.raises(ic.ErroPlanilha, match='cabeçalho'):
        ic.ler_geral(caminho)


def test_arquivo_que_nao_e_planilha(tmp_path):
    caminho = tmp_path / 'x.xlsx'
    caminho.write_bytes(b'nao e xlsx')
    with pytest.raises(ic.ErroPlanilha):
        ic.ler_geral(str(caminho))


@pytest.mark.parametrize('mes_planilha, mes, divergente', [
    ('ABRIL  2026', '2026-06', True),
    ('JUNHO 2026', '2026-06', False),
    ('Março / 2026', '2026-03', False),
    ('JUNHO 2025', '2026-06', True),
    (None, '2026-06', False),
])
def test_mes_divergente(mes_planilha, mes, divergente):
    assert ic._mes_divergente(mes_planilha, mes) is divergente


@pytest.mark.skipif(not os.path.exists(PLANILHA_REAL), reason='planilha de junho/2026 não está em dados/')
def test_planilha_real_de_junho():
    d = ic.ler_geral(PLANILHA_REAL)
    custos = [i for i in d['itens'] if i['categoria'] == 'custo']
    assert len(d['unidades']) == 45
    assert len({u['grupo'] for u in d['unidades']}) == 3
    assert len(custos) == 58
    assert len([i for i in d['itens'] if i['categoria'] == 'rateio']) == 11
    assert {i['tipo'] for i in custos} == {'Pessoal', 'Material de Consumo', 'Despesas Gerais', 'Custos Indiretos'}
    assert sum(i['total'] or 0 for i in custos) == pytest.approx(d['custo_total_planilha'], abs=0.01)
    assert d['avisos'] == []


def test_envio_exige_mes_de_referencia(cliente, logar, admin_id, tmp_path):
    logar(admin_id)
    with open(_planilha(tmp_path, LINHAS), 'rb') as f:
        r = cliente.post('/admin/custos/importar-geral', content_type='multipart/form-data',
                         data={'csrf': 'tok', 'mes': '', 'planilha': (io.BytesIO(f.read()), 'p.xlsx')})
    assert r.status_code == 400
    assert 'Informe o mês de referência da planilha' in r.get_data(as_text=True)


def test_envio_sem_banco_mostra_erro_na_janela(cliente, logar, admin_id, tmp_path):
    logar(admin_id)
    caminho = _planilha(tmp_path, LINHAS)
    with open(caminho, 'rb') as f:
        r = cliente.post('/admin/custos/importar-geral', content_type='multipart/form-data',
                         data={'csrf': 'tok', 'mes': '2026-06', 'planilha': (io.BytesIO(f.read()), 'p.xlsx')})
    assert r.status_code == 400
    assert 'Não foi possível preparar a importação' in r.get_data(as_text=True)
